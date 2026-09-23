from pathlib import Path

import docx
from fpdf import FPDF

from analyzer.contracts.text import from_docx, from_pdf, from_plain

FIXTURE = Path(__file__).parent / "fixtures" / "contract_glow.txt"


def _pdf(path: Path, text: str | None) -> Path:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    if text:
        pdf.multi_cell(0, 5, text.encode("latin-1", "replace").decode("latin-1"))
    else:
        pdf.rect(20, 20, 100, 60, style="F")  # an "image-only" page
    pdf.output(str(path))
    return path


def test_pdf_text_is_extracted(tmp_path):
    ct = from_pdf(_pdf(tmp_path / "c.pdf", FIXTURE.read_text()))
    assert "INFLUENCER SERVICES AGREEMENT" in ct.text
    assert "ninety (90) days" in ct.text
    assert not ct.scanned


def test_pdf_without_text_is_flagged_as_scanned(tmp_path):
    ct = from_pdf(_pdf(tmp_path / "scan.pdf", None))
    assert ct.scanned
    assert ct.text == ""


def test_docx_paragraphs_and_tables(tmp_path):
    d = docx.Document()
    d.add_paragraph("Brand will pay a fee of £2,500.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Usage"
    table.rows[0].cells[1].text = "12 months"
    path = tmp_path / "c.docx"
    d.save(str(path))
    ct = from_docx(path)
    assert "£2,500" in ct.text
    assert "Usage | 12 months" in ct.text


def test_plain_text_is_tidied():
    ct = from_plain("A\r\n\r\n\r\n\r\nB  \n")
    assert ct.text == "A\n\nB"

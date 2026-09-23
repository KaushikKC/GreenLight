"""Plain text from contract files. Scanned PDFs are detected, not OCR'd:
they go to the LLM as a document instead (BUILD_PLAN §7.2)."""

import re
from dataclasses import dataclass
from pathlib import Path

import docx
import pdfplumber

# Below this many characters per page we assume the PDF is a scan.
SCANNED_CHARS_PER_PAGE = 50


@dataclass
class ContractText:
    text: str
    pages: int
    scanned: bool


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def from_pdf(path: Path) -> ContractText:
    with pdfplumber.open(path) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    text = _tidy("\n\n".join(pages))
    n = max(len(pages), 1)
    return ContractText(text=text, pages=len(pages), scanned=len(text) < SCANNED_CHARS_PER_PAGE * n)


def from_docx(path: Path) -> ContractText:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ContractText(text=_tidy("\n".join(parts)), pages=1, scanned=False)


def from_plain(text: str) -> ContractText:
    return ContractText(text=_tidy(text), pages=1, scanned=False)
